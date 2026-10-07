"""Report generation: multi-sheet Excel workbook (openpyxl) or CSV set.

The "suggested command" columns, and the Inventory change-request columns
(remove / delete / rollback), are informational text for the human running
the change control; this tool never executes anything.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import commands
from .analyzer import AnalysisResult, ObjectVerdict, Verdict
from .correlator import Correlation
from .models import Node, Pool, PoolMember
from .parsing import ParsedData

VERDICT_FILLS = {
    Verdict.ORPHAN: "FFC7CE",  # red
    Verdict.MANUAL_REVIEW: "FFEB9C",  # yellow
    Verdict.UNRELIABLE_STANDBY: "FFEB9C",  # yellow
    Verdict.UNRELIABLE_INVENTORY: "FFEB9C",
    Verdict.INACTIVE: "FCD5B4",  # orange
    Verdict.OFFLINE_CANDIDATE: "CCC0DA",  # purple
    Verdict.IN_USE: "C6EFCE",  # green
}

HEADER_FILL = "D9D9D9"
MAX_COLUMN_WIDTH = 60

# Written by the opt-in `f5audit ping` post-process. The Inventory sheet
# reserves these two columns up front so the change-request columns after
# them keep a fixed position whether or not ping is ever run.
PING_STATUS_HEADER = "Ping (from F5)"
PING_NOTE_HEADER = "Ping note"

CHANGE_HEADERS = [
    "Remove node from pool",
    "Delete node",
    "Create node (rollback)",
    "Add node back to pool (rollback)",
]

# Only these verdicts get change and rollback commands; every other one
# (MANUAL REVIEW, INACTIVE, UNRELIABLE) stays blank, the conservative choice.
COMMAND_VERDICTS = (Verdict.ORPHAN, Verdict.OFFLINE_CANDIDATE)

NODE_COMMAND_HEADERS = [
    "Remove node from pool(s)",
    "Delete node",
    "Create node (rollback)",
    "Add node back to pool(s) (rollback)",
]
POOL_COMMAND_HEADERS = [
    "Detach pool from virtual servers",
    "Delete pool",
    "Recreate pool (rollback)",
    "Reattach pool to virtual servers (rollback)",
]
VIRTUAL_COMMAND_HEADERS = [
    "Delete virtual server",
    "Recreate virtual server (rollback)",
]
CHAIN_COMMAND_HEADERS = [
    "Delete virtual servers",
    "Delete pool",
    "Delete nodes",
    "Recreate nodes (rollback)",
    "Recreate pool (rollback)",
    "Recreate virtual servers (rollback)",
]

VIRTUAL_ADDRESS_NOTE = (
    "Deleting the last virtual server on an address also deletes its "
    "virtual-address (auto-delete); re-creating the virtual server brings "
    "it back with default settings, so record any non-default ARP, ICMP, "
    "route-advertisement or traffic-group setting before the change."
)
PROFILES_MISSING_NOTE = (
    "No rollback: the virtual server's profiles were not collected; re-run "
    "'f5audit collect --save-raw' on this raw directory to backfill them."
)


@dataclass
class ReportTable:
    title: str
    headers: list[str]
    rows: list[list[object]] = field(default_factory=list)
    verdict_columns: tuple[int, ...] = ()  # 0-based indexes into headers


def default_report_name(hostname: str, fmt: str = "xlsx") -> str:
    safe_host = re.sub(r"[^A-Za-z0-9._-]", "_", hostname or "unknown")
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    suffix = "" if fmt == "csv" else ".xlsx"
    return f"f5audit_{safe_host}_{stamp}{suffix}"


# ---------------------------------------------------------------------------
# Table construction
# ---------------------------------------------------------------------------


def _join(values) -> str:
    return ", ".join(sorted(values)) if values else ""


def _network_cells(parsed: ParsedData, node) -> list[object]:
    """[ARP MAC, Network note] appended at the end of node rows.

    Informational only: ARP absence means "not seen at collection time",
    and on a standby unit the table reflects that unit, not the pair.
    """
    if node is None:
        return ["", ""]
    if not parsed.network_collected:
        return ["", "network data not collected"]
    info = parsed.network.get(node.address)
    if info is None:
        return ["", ""]
    note = info.connectivity
    if note and parsed.system.failover_state == "standby":
        note += " [standby unit: ARP reflects this unit only]"
    return [info.arp_mac, note]


def _change_cells(
    node: Node | None,
    node_verdict: ObjectVerdict | None,
    pool_path: str | None,
    member: PoolMember | None,
) -> list[str]:
    """[remove from pool, delete node, create node, add back to pool].

    Only OFFLINE and ORPHAN nodes get commands; every other verdict
    (including MANUAL REVIEW) stays blank, the conservative choice.
    """
    blank = ["", "", "", ""]
    if node is None or node_verdict is None:
        return blank
    if not _has_commands(node_verdict):
        return blank
    delete = commands.delete_node_command(node.full_path)
    create = commands.node_create_command(node)
    if node_verdict.verdict == Verdict.ORPHAN or member is None or not pool_path:
        return ["", delete, create, ""]
    remove = commands.remove_member_command(pool_path, member)
    add = commands.add_member_command(pool_path, member)
    return [remove, delete, create, add]


def _has_commands(verdict: ObjectVerdict | None) -> bool:
    return verdict is not None and verdict.verdict in COMMAND_VERDICTS


def _memberships(parsed: ParsedData, node_path: str) -> list[tuple[str, PoolMember]]:
    return [
        (pool_path, member)
        for pool_path, pool in sorted(parsed.pools.items())
        for member in pool.members
        if member.node_full_path == node_path
    ]


def _node_command_cells(parsed: ParsedData, node: Node, verdict: ObjectVerdict) -> list[str]:
    """NODE_COMMAND_HEADERS cells: one remove/add line per pool membership."""
    if not _has_commands(verdict):
        return ["", "", "", ""]
    memberships = _memberships(parsed, node.full_path)
    return [
        "\n".join(
            commands.remove_member_command(pool_path, member) for pool_path, member in memberships
        ),
        commands.delete_node_command(node.full_path),
        commands.node_create_command(node),
        "\n".join(
            commands.add_member_command(pool_path, member) for pool_path, member in memberships
        ),
    ]


def _with_note(notes: str, extra: str) -> str:
    return f"{notes} {extra}".strip() if extra else notes


def _virtual_command_cells(virtual, verdict: ObjectVerdict) -> tuple[list[str], str]:
    """(VIRTUAL_COMMAND_HEADERS cells, extra note)."""
    if not _has_commands(verdict):
        return ["", ""], ""
    note = VIRTUAL_ADDRESS_NOTE
    if not virtual.profiles_collected:
        note += " " + PROFILES_MISSING_NOTE
    return [
        commands.delete_virtual_command(virtual.full_path),
        commands.virtual_create_command(virtual),
    ], note


def build_tables(
    parsed: ParsedData, correlation: Correlation, analysis: AnalysisResult
) -> dict[str, ReportTable]:
    return {
        "summary": _build_summary(parsed, analysis),
        "inventory": _build_inventory(parsed, correlation, analysis),
        "orphan_nodes": _build_orphan_nodes(parsed, analysis),
        "pools": _build_pools(parsed, correlation, analysis),
        "inactive_virtuals": _build_inactive_virtuals(parsed, analysis),
        "dead_chains": _build_dead_chains(parsed, correlation, analysis),
        "orphan_monitors": _build_orphan_monitors(parsed, analysis),
        "manual_review": _build_manual_review(analysis),
    }


def _build_summary(parsed: ParsedData, analysis: AnalysisResult) -> ReportTable:
    system = parsed.system
    table = ReportTable("Summary", ["Item", "Value"])
    table.rows = [
        ["Hostname", system.hostname],
        ["BIG-IP version", system.version],
        ["HA state", system.failover_state],
        ["Active device", system.active_device],
        ["Uptime context", system.uptime],
        ["Collection timestamp (UTC)", system.collection_timestamp],
        ["Partitions collected", _join(system.partitions_collected)],
        ["Partitions denied", _join(system.partitions_denied)],
        ["Missing endpoints", _join(system.missing_endpoints)],
        ["Traffic analysis skipped", "YES" if analysis.stats_analysis_skipped else "no"],
        [
            "Network data (ARP/self-IP)",
            "collected" if parsed.network_collected else "not collected (old cache)",
        ],
        ["", ""],
        ["Objects", ""],
        ["Nodes", len(parsed.nodes)],
        ["Pools", len(parsed.pools)],
        ["Virtual servers", len(parsed.virtuals)],
        ["iRules", len(parsed.irules)],
        ["Policies", len(parsed.policies)],
        ["Monitors", len(parsed.monitors)],
        ["", ""],
        ["Verdict counts", ""],
    ]
    if system.resumed_at:
        table.rows.insert(
            6, ["Collection resumed (mixed timestamps)", ", ".join(system.resumed_at)]
        )
    for verdict, count in sorted(analysis.verdict_counts().items()):
        table.rows.append([verdict, count])
    if analysis.warnings:
        table.rows.append(["", ""])
        table.rows.append(["WARNINGS", ""])
        for warning in analysis.warnings:
            table.rows.append(["!", warning])
    return table


def _vs_summary(parsed: ParsedData, vs_paths) -> tuple:
    """(names, destinations, states, total_conns, attached iRules) joined
    for a VS set."""
    names, destinations, states, conns, irules = [], [], [], [], []
    for path in sorted(vs_paths):
        virtual = parsed.virtuals.get(path)
        if not virtual:
            names.append(path)
            continue
        names.append(path)
        destinations.append(virtual.destination)
        states.append(f"{virtual.admin_state}/{virtual.availability or '?'}")
        if virtual.total_conns is not None:
            conns.append(str(virtual.total_conns))
        irules.extend(irule for irule in virtual.irules if irule not in irules)
    return (
        ", ".join(names),
        ", ".join(destinations),
        ", ".join(states),
        ", ".join(conns),
        ", ".join(irules),
    )


def _build_inventory(
    parsed: ParsedData, correlation: Correlation, analysis: AnalysisResult
) -> ReportTable:
    headers = [
        "Node",
        "IP",
        "Partition",
        "Node status",
        "Node verdict",
        "Pool",
        "Port",
        "Member status",
        "Effective monitor",
        "LB method",
        "Virtual servers",
        "VIP:Port",
        "VS status",
        "VS total conns",
        "VS iRules",
        "iRules selecting pool",
        "Policies forwarding to pool",
        "Pool verdict",
        "Pool notes",
        "ARP MAC",
        "Network note",
        PING_STATUS_HEADER,
        PING_NOTE_HEADER,
        *CHANGE_HEADERS,
    ]
    # Both verdict columns are colored independently: node verdict at 4,
    # pool verdict at 17.
    table = ReportTable("Inventory", headers, verdict_columns=(4, 17))

    nodes_in_pools = set()
    for pool_path, pool in sorted(parsed.pools.items()):
        verdict = analysis.pool_verdicts.get(pool_path)
        vs_paths = correlation.pool_to_virtuals.get(pool_path, set())
        vs_names, vips, vs_states, vs_conns, vs_irules = _vs_summary(parsed, vs_paths)
        irule_refs = _join(correlation.pool_to_irules.get(pool_path, set()))
        policy_refs = _join(correlation.pool_to_policies.get(pool_path, set()))
        members = pool.members or [None]
        for member in members:
            node = parsed.nodes.get(member.node_full_path) if member else None
            node_verdict = analysis.node_verdicts.get(member.node_full_path) if member else None
            if member:
                nodes_in_pools.add(member.node_full_path)
            monitor = _join(pool.monitors) or (node.monitor if node else "")
            table.rows.append(
                [
                    member.node_full_path if member else "(no members)",
                    node.address if node else "",
                    pool.partition,
                    f"{node.admin_state}/{node.availability or '?'}" if node else "",
                    node_verdict.verdict if node_verdict else "",
                    pool_path,
                    member.port if member else "",
                    f"{member.admin_state}/{member.availability or '?'}" if member else "",
                    monitor,
                    pool.lb_method,
                    vs_names,
                    vips,
                    vs_states,
                    vs_conns,
                    vs_irules,
                    irule_refs,
                    policy_refs,
                    verdict.verdict if verdict else "",
                    verdict.notes if verdict else "",
                ]
                + _network_cells(parsed, node)
                + ["", ""]
                + _change_cells(node, node_verdict, pool_path, member)
            )

    # Nodes that belong to no pool get their own rows.
    for node_path, node in sorted(parsed.nodes.items()):
        if node_path in nodes_in_pools:
            continue
        verdict = analysis.node_verdicts.get(node_path)
        table.rows.append(
            [
                node_path,
                node.address,
                node.partition,
                f"{node.admin_state}/{node.availability or '?'}",
                verdict.verdict if verdict else "",
                "",
                "",
                "",
                node.monitor,
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                # No pool: mirror the node verdict into the pool columns so a
                # standalone node still shows a verdict here too.
                verdict.verdict if verdict else "",
                verdict.notes if verdict else "",
            ]
            + _network_cells(parsed, node)
            + ["", ""]
            + _change_cells(node, verdict, None, None)
        )
    return table


def _build_orphan_nodes(parsed: ParsedData, analysis: AnalysisResult) -> ReportTable:
    headers = [
        "Node",
        "IP",
        "Partition",
        "Monitor",
        "Verdict",
        "Notes",
        *NODE_COMMAND_HEADERS,
        "ARP MAC",
        "Network note",
    ]
    table = ReportTable("Orphan Nodes", headers, verdict_columns=(4,))
    for path, verdict in sorted(analysis.node_verdicts.items()):
        if verdict.verdict == Verdict.IN_USE:
            continue
        node = parsed.nodes[path]
        table.rows.append(
            [
                path,
                node.address,
                node.partition,
                node.monitor,
                verdict.verdict,
                verdict.notes,
                *_node_command_cells(parsed, node, verdict),
            ]
            + _network_cells(parsed, node)
        )
    return table


def _build_pools(
    parsed: ParsedData, correlation: Correlation, analysis: AnalysisResult
) -> ReportTable:
    headers = [
        "Pool",
        "Partition",
        "LB method",
        "Members",
        "Monitors",
        "Virtual servers",
        "iRules selecting pool",
        "Policies forwarding to pool",
        "Verdict",
        "Notes",
        *POOL_COMMAND_HEADERS,
    ]
    table = ReportTable("Orphan-Inactive Pools", headers, verdict_columns=(8,))
    for path, verdict in sorted(analysis.pool_verdicts.items()):
        if verdict.verdict == Verdict.IN_USE:
            continue
        pool = parsed.pools[path]
        cells = _pool_command_cells(correlation, pool, verdict)
        table.rows.append(
            [
                path,
                pool.partition,
                pool.lb_method,
                ", ".join(f"{m.node_full_path}:{m.port}" for m in pool.members),
                _join(pool.monitors),
                _join(correlation.pool_to_virtuals.get(path, set())),
                _join(correlation.pool_to_irules.get(path, set())),
                _join(correlation.pool_to_policies.get(path, set())),
                verdict.verdict,
                verdict.notes,
                *cells,
            ]
        )
    return table


def _pool_command_cells(correlation: Correlation, pool: Pool, verdict: ObjectVerdict) -> list[str]:
    """POOL_COMMAND_HEADERS cells. Detach covers default pools, the only
    reference a pool with a command verdict can have: one named by an iRule
    or a policy is never ORPHAN, and MANUAL REVIEW when dead."""
    if not _has_commands(verdict):
        return ["", "", "", ""]
    virtuals = sorted(correlation.pool_to_virtuals.get(pool.full_path, set()))
    return [
        "\n".join(commands.detach_pool_command(vs) for vs in virtuals),
        commands.delete_pool_command(pool.full_path),
        commands.pool_create_command(pool),
        "\n".join(commands.reattach_pool_command(vs, pool.full_path) for vs in virtuals),
    ]


def _build_inactive_virtuals(parsed: ParsedData, analysis: AnalysisResult) -> ReportTable:
    headers = [
        "Virtual server",
        "VIP:Port",
        "Default pool",
        "iRules",
        "State",
        "Availability",
        "Total conns",
        "Bits in",
        "Bits out",
        "Verdict",
        "Notes",
        *VIRTUAL_COMMAND_HEADERS,
    ]
    table = ReportTable("Inactive Virtual Servers", headers, verdict_columns=(9,))
    for path, verdict in sorted(analysis.virtual_verdicts.items()):
        if verdict.verdict == Verdict.IN_USE:
            continue
        virtual = parsed.virtuals[path]
        cells, note = _virtual_command_cells(virtual, verdict)
        table.rows.append(
            [
                path,
                virtual.destination,
                virtual.default_pool,
                _join(virtual.irules),
                virtual.admin_state,
                virtual.availability,
                virtual.total_conns,
                virtual.bits_in,
                virtual.bits_out,
                verdict.verdict,
                _with_note(verdict.notes, note),
                *cells,
            ]
        )
    return table


def _build_dead_chains(
    parsed: ParsedData, correlation: Correlation, analysis: AnalysisResult
) -> ReportTable:
    """One row per dead chain, grouped by pool: the artifact to take to the
    config owner. Commands are informational text, one column per step:
    VS -> pool -> nodes, then the rollback in reverse. Only OFFLINE objects
    get commands, and only when tmsh would accept them in this order: a
    pool is deleted only if nothing outside this row still references it,
    and a node only if this pool is deleted here and was its last pool."""
    headers = [
        "Pool",
        "Partition",
        "Pool status",
        "Members",
        "Member statuses",
        "Node verdicts",
        "Virtual servers",
        "VS statuses",
        "VS verdicts",
        "Verdict",
        "Notes",
        *CHAIN_COMMAND_HEADERS,
    ]
    table = ReportTable("Dead Chains", headers, verdict_columns=(9,))
    for path in sorted(analysis.offline_pools):
        verdict = analysis.pool_verdicts[path]
        pool = parsed.pools[path]
        vs_paths = sorted(vs for vs, pools in correlation.virtual_to_pools.items() if path in pools)
        node_paths = sorted({member.node_full_path for member in pool.members})
        notes = []

        vs_states, vs_verdict_labels, deleted_virtuals = [], [], []
        for vs_path in vs_paths:
            virtual = parsed.virtuals.get(vs_path)
            vs_states.append(
                f"{virtual.admin_state}/{virtual.availability or '?'}" if virtual else "?"
            )
            vs_verdict = analysis.virtual_verdicts.get(vs_path)
            vs_verdict_labels.append(vs_verdict.verdict if vs_verdict else "")
            if virtual and _has_commands(vs_verdict):
                deleted_virtuals.append(virtual)
        if deleted_virtuals:
            notes.append(VIRTUAL_ADDRESS_NOTE)
            if any(not virtual.profiles_collected for virtual in deleted_virtuals):
                notes.append(PROFILES_MISSING_NOTE)

        pool_deleted = False
        if _has_commands(verdict):
            deleted_paths = {virtual.full_path for virtual in deleted_virtuals}
            kept = sorted(correlation.pool_to_virtuals.get(path, set()) - deleted_paths)
            if kept:
                notes.append(
                    f"Pool not deleted: still the default pool of {', '.join(kept)}, "
                    "which is not a removal candidate (see Orphan-Inactive Pools)."
                )
            else:
                pool_deleted = True

        node_verdict_labels, deleted_nodes = [], []
        for node_path in node_paths:
            node_verdict = analysis.node_verdicts.get(node_path)
            node_verdict_labels.append(node_verdict.verdict if node_verdict else "")
            node = parsed.nodes.get(node_path)
            if node is None or not _has_commands(node_verdict):
                continue
            if pool_deleted and correlation.node_to_pools.get(node_path, set()) <= {path}:
                deleted_nodes.append(node)
            else:
                notes.append(
                    f"Node {node_path}: remove it from its pools first (see Orphan Nodes)."
                )

        table.rows.append(
            [
                path,
                pool.partition,
                pool.availability,
                ", ".join(f"{m.node_full_path}:{m.port}" for m in pool.members),
                ", ".join(f"{m.admin_state}/{m.availability or '?'}" for m in pool.members),
                ", ".join(node_verdict_labels),
                ", ".join(vs_paths),
                ", ".join(vs_states),
                ", ".join(vs_verdict_labels),
                verdict.verdict,
                " ".join([verdict.notes, *notes]).strip(),
                "\n".join(
                    commands.delete_virtual_command(virtual.full_path)
                    for virtual in deleted_virtuals
                ),
                commands.delete_pool_command(path) if pool_deleted else "",
                "\n".join(commands.delete_node_command(node.full_path) for node in deleted_nodes),
                "\n".join(commands.node_create_command(node) for node in deleted_nodes),
                commands.pool_create_command(pool) if pool_deleted else "",
                "\n".join(
                    command
                    for command in map(commands.virtual_create_command, deleted_virtuals)
                    if command
                ),
            ]
        )
    return table


def _build_orphan_monitors(parsed: ParsedData, analysis: AnalysisResult) -> ReportTable:
    headers = [
        "Monitor",
        "Type",
        "Partition",
        "Verdict",
        "Notes",
        "Suggested command (informational)",
    ]
    table = ReportTable("Orphan Monitors", headers, verdict_columns=(3,))
    for path, verdict in sorted(analysis.monitor_verdicts.items()):
        if verdict.verdict == Verdict.IN_USE:
            continue
        monitor = parsed.monitors[path]
        table.rows.append(
            [
                path,
                monitor.type,
                monitor.partition,
                verdict.verdict,
                verdict.notes,
                f"delete ltm monitor {monitor.type} {path}"
                if verdict.verdict == Verdict.ORPHAN
                else "",
            ]
        )
    return table


def _build_manual_review(analysis: AnalysisResult) -> ReportTable:
    headers = ["Object type", "Object", "Reason", "Caused by"]
    table = ReportTable("Manual Review", headers)
    for item in analysis.manual_review:
        table.rows.append([item.object_type, item.full_path, item.reason, item.caused_by])
    return table


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------


def write_xlsx(tables: dict[str, ReportTable], out_path: str) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)

    for table in tables.values():
        sheet = workbook.create_sheet(title=table.title[:31])
        sheet.append(table.headers)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in table.rows:
            sheet.append(row)
            for column in table.verdict_columns:
                verdict = row[column]
                color = VERDICT_FILLS.get(str(verdict))
                if color:
                    cell = sheet.cell(row=sheet.max_row, column=column + 1)
                    cell.fill = PatternFill("solid", fgColor=color)
        sheet.freeze_panes = "A2"
        last_column = get_column_letter(len(table.headers))
        sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"
        _autofit_columns(sheet, table)
    workbook.save(out_path)


def _autofit_columns(sheet, table: ReportTable) -> None:
    for index, header in enumerate(table.headers, start=1):
        width = len(str(header))
        for row in table.rows:
            value = row[index - 1]
            if value is not None:
                width = max(width, len(str(value)))
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, MAX_COLUMN_WIDTH)


def write_csv(tables: dict[str, ReportTable], out_dir: str) -> list[str]:
    directory = Path(out_dir)
    directory.mkdir(parents=True, exist_ok=True)
    written = []
    for key, table in tables.items():
        path = directory / f"{key}.csv"
        with open(path, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(table.headers)
            writer.writerows(table.rows)
        written.append(str(path))
    return written

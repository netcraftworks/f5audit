"""Report tests: table content, Excel structure, CSV output."""

import csv

from openpyxl import load_workbook

from f5audit.analyzer import Analyzer, Verdict
from f5audit.correlator import correlate
from f5audit.models import IRule, Node, PoolMember
from f5audit.parsing import parse_collection
from f5audit.report import (
    POOL_COMMAND_HEADERS,
    build_tables,
    default_report_name,
    write_csv,
    write_xlsx,
)
from tests.conftest import build_collection

# Inventory layout: network columns T:U, reserved ping columns V:W, and the
# change-request columns X:AA the user copies into the change request.
NETWORK_COLUMNS = slice(19, 21)
PING_COLUMNS = slice(21, 23)
CHANGE_COLUMNS = slice(23, 27)


def make_tables():
    parsed = parse_collection(build_collection())
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    return parsed, build_tables(parsed, correlation, analysis)


def test_expected_sheets_exist():
    _, tables = make_tables()
    assert list(tables) == [
        "summary",
        "inventory",
        "orphan_nodes",
        "pools",
        "inactive_virtuals",
        "dead_chains",
        "orphan_monitors",
        "manual_review",
    ]


def test_inventory_has_member_rows_and_orphan_node_rows():
    _, tables = make_tables()
    inventory = tables["inventory"]
    first_column = [row[0] for row in inventory.rows]
    assert "/Common/node-web-1" in first_column  # pool member row
    assert "/Common/node-orphan" in first_column  # node without pool

    member_row = next(
        r for r in inventory.rows if r[0] == "/Common/node-web-1" and r[5] == "/Common/pool-web"
    )
    assert member_row[1] == "10.0.0.1"
    assert member_row[4] == Verdict.IN_USE  # node verdict
    assert member_row[6] == "80"
    assert "/Common/vs-web" in member_row[10]
    assert member_row[14] == "/Common/irule-static"  # iRules attached to the VS
    assert member_row[17] == Verdict.IN_USE  # pool verdict

    orphan_row = next(r for r in inventory.rows if r[0] == "/Common/node-orphan")
    assert orphan_row[4] == Verdict.ORPHAN
    assert orphan_row[17] == Verdict.ORPHAN


def test_orphan_sheets_only_contain_non_in_use_objects():
    _, tables = make_tables()
    assert [row[0] for row in tables["orphan_nodes"].rows] == [
        "/Common/node-dead",
        "/Common/node-orphan",
    ]
    pool_names = [row[0] for row in tables["pools"].rows]
    assert "/Common/pool-orphan" in pool_names
    assert "/Common/pool-idle" in pool_names
    assert "/Common/pool-dead" in pool_names
    assert "/Common/pool-web" not in pool_names
    monitor_names = [row[0] for row in tables["orphan_monitors"].rows]
    assert monitor_names == ["/Common/mon-orphan"]


def cells(table, row):
    return dict(zip(table.headers, row))


def table_row(tables, key, first_cell):
    table = tables[key]
    return cells(table, next(r for r in table.rows if r[0] == first_cell))


def test_orphan_nodes_commands_for_offline_and_orphan():
    _, tables = make_tables()
    dead = table_row(tables, "orphan_nodes", "/Common/node-dead")
    assert dead["Remove node from pool(s)"] == (
        "modify ltm pool /Common/pool-dead members delete { /Common/node-dead:443 }"
    )
    assert dead["Delete node"] == "delete ltm node /Common/node-dead"
    assert dead["Create node (rollback)"] == (
        "create ltm node /Common/node-dead address 10.0.0.50 monitor /Common/icmp"
    )
    assert dead["Add node back to pool(s) (rollback)"] == (
        "modify ltm pool /Common/pool-dead members add { /Common/node-dead:443 }"
    )
    orphan = table_row(tables, "orphan_nodes", "/Common/node-orphan")
    assert orphan["Remove node from pool(s)"] == ""
    assert orphan["Delete node"] == "delete ltm node /Common/node-orphan"
    assert orphan["Add node back to pool(s) (rollback)"] == ""


def test_pool_commands_only_for_orphan_and_offline():
    _, tables = make_tables()
    offline = table_row(tables, "pools", "/Common/pool-dead")
    assert offline["Detach pool from virtual servers"] == (
        "modify ltm virtual /Common/vs-dead pool none"
    )
    assert offline["Delete pool"] == "delete ltm pool /Common/pool-dead"
    assert offline["Recreate pool (rollback)"] == (
        "create ltm pool /Common/pool-dead members add { /Common/node-dead:443 } "
        "monitor /Common/mon-used"
    )
    assert offline["Reattach pool to virtual servers (rollback)"] == (
        "modify ltm virtual /Common/vs-dead pool /Common/pool-dead"
    )
    orphan = table_row(tables, "pools", "/Common/pool-orphan")
    assert orphan["Detach pool from virtual servers"] == ""
    assert orphan["Delete pool"] == "delete ltm pool /Common/pool-orphan"
    assert (
        orphan["Recreate pool (rollback)"]
        == "create ltm pool /Common/pool-orphan monitor /Common/tcp"
    )
    inactive = table_row(tables, "pools", "/Common/pool-idle")
    assert inactive["Delete pool"] == ""
    assert inactive["Recreate pool (rollback)"] == ""


def test_pool_commands_withheld_when_an_irule_or_policy_references_it():
    parsed = parse_collection(build_collection())
    parsed.pools["/Common/pool-irule"].members = []  # empty: OFFLINE via vs-web's iRule
    tables = analyze(parsed)
    row = table_row(tables, "pools", "/Common/pool-irule")
    assert row["Verdict"] == Verdict.OFFLINE_CANDIDATE
    assert [row[h] for h in POOL_COMMAND_HEADERS] == ["", "", "", ""]
    assert "Commands withheld: referenced by iRule(s) /Common/irule-static" in row["Notes"]


def test_virtual_commands_for_offline_virtual():
    _, tables = make_tables()
    row = table_row(tables, "inactive_virtuals", "/Common/vs-dead")
    assert row["Delete virtual server"] == "delete ltm virtual /Common/vs-dead"
    assert row["Recreate virtual server (rollback)"] == (
        "create ltm virtual /Common/vs-dead destination /Common/192.0.2.13:443 "
        "mask 255.255.255.255 ip-protocol tcp pool /Common/pool-dead "
        "profiles add { /Common/tcp { context all } /Common/http { context all } "
        "/Common/clientssl-example { context clientside } } "
        "persist replace-all-with { /Common/cookie { default yes } } "
        "fallback-persistence /Common/source_addr "
        "source-address-translation { type snat pool /Common/snat-example } "
        "vlans-enabled vlans add { /Common/vlan-external } "
        'description "Decommissioned app example.net"'
    )
    assert "virtual-address" in row["Notes"]
    idle = table_row(tables, "inactive_virtuals", "/Common/vs-idle")
    assert idle["Verdict"] == Verdict.INACTIVE
    assert idle["Delete virtual server"] == ""


def test_virtual_rollback_withheld_without_collected_profiles():
    collection = build_collection()
    del collection.datasets["ltm_virtual_profiles@/Common/vs-dead"]
    tables = analyze(parse_collection(collection))
    row = table_row(tables, "inactive_virtuals", "/Common/vs-dead")
    assert row["Delete virtual server"] == "delete ltm virtual /Common/vs-dead"
    assert row["Recreate virtual server (rollback)"] == ""
    assert "profiles were not collected" in row["Notes"]


def test_dead_chains_sheet_groups_the_whole_chain():
    _, tables = make_tables()
    assert [row[0] for row in tables["dead_chains"].rows] == ["/Common/pool-dead"]
    row = table_row(tables, "dead_chains", "/Common/pool-dead")
    assert row["Members"] == "/Common/node-dead:443"
    assert row["Virtual servers"] == "/Common/vs-dead"
    assert row["Verdict"] == Verdict.OFFLINE_CANDIDATE
    assert row["Delete virtual servers"] == "delete ltm virtual /Common/vs-dead"
    assert row["Delete pool"] == "delete ltm pool /Common/pool-dead"
    assert row["Delete nodes"] == "delete ltm node /Common/node-dead"
    assert row["Recreate nodes (rollback)"] == (
        "create ltm node /Common/node-dead address 10.0.0.50 monitor /Common/icmp"
    )
    assert row["Recreate pool (rollback)"].startswith("create ltm pool /Common/pool-dead ")
    assert row["Recreate virtual servers (rollback)"].startswith(
        "create ltm virtual /Common/vs-dead "
    )


def test_dead_chains_sheet_omits_node_command_when_alive_elsewhere():
    parsed = parse_collection(build_collection())
    # node-dead is also an available member of pool-web.
    parsed.pools["/Common/pool-web"].members.append(
        PoolMember(
            node_full_path="/Common/node-dead",
            port="80",
            partition="Common",
            admin_state="monitor-enabled",
            availability="available",
        )
    )
    row = table_row(analyze(parsed), "dead_chains", "/Common/pool-dead")
    assert row["Delete nodes"] == ""
    assert row["Recreate nodes (rollback)"] == ""
    # The IN USE node holds the pool and its virtual server at MANUAL
    # REVIEW: the chain is listed, but nothing in it gets a command.
    assert row["Node verdicts"] == Verdict.IN_USE
    assert row["Verdict"] == Verdict.MANUAL_REVIEW
    assert row["VS verdicts"] == Verdict.MANUAL_REVIEW
    assert row["Delete pool"] == ""
    assert row["Delete virtual servers"] == ""
    assert "still IN USE" in row["Notes"]


def test_dead_chains_sheet_keeps_capped_pool_without_pool_or_node_command():
    parsed = parse_collection(build_collection())
    parsed.irules["/Common/irule-dyn"] = IRule(
        full_path="/Common/irule-dyn",
        partition="Common",
        name="irule-dyn",
        definition="pool $x",
        has_dynamic_pool_selection=True,
    )
    parsed.virtuals["/Common/vs-web"].irules.append("/Common/irule-dyn")
    tables = analyze(parsed)
    assert [row[0] for row in tables["dead_chains"].rows] == ["/Common/pool-dead"]
    row = table_row(tables, "dead_chains", "/Common/pool-dead")
    assert row["Verdict"] == Verdict.MANUAL_REVIEW
    assert row["Node verdicts"] == Verdict.OFFLINE_CANDIDATE
    assert row["Delete virtual servers"] == "delete ltm virtual /Common/vs-dead"
    assert row["Delete pool"] == ""
    # The node is still a member of the pool that stays: tmsh would refuse
    # the delete, so the Orphan Nodes sheet carries its commands instead.
    assert row["Delete nodes"] == ""
    assert "remove it from its pools first" in row["Notes"]


def test_dead_chains_withholds_pool_delete_while_a_kept_virtual_uses_it():
    parsed = parse_collection(build_collection())
    parsed.virtuals["/Common/vs-dead"].app_service = "/Common/adfs.app/adfs"
    row = table_row(analyze(parsed), "dead_chains", "/Common/pool-dead")
    assert row["VS verdicts"] == Verdict.MANUAL_REVIEW
    assert row["Delete virtual servers"] == ""
    assert row["Delete pool"] == ""
    assert "still the default pool of /Common/vs-dead" in row["Notes"]


def test_dead_chains_includes_empty_pool_chain():
    parsed = parse_collection(build_collection())
    parsed.pools["/Common/pool-idle"].members = []
    row = table_row(analyze(parsed), "dead_chains", "/Common/pool-idle")
    assert row["Verdict"] == Verdict.OFFLINE_CANDIDATE
    assert row["Delete virtual servers"] == "delete ltm virtual /Common/vs-idle"
    assert row["Delete pool"] == "delete ltm pool /Common/pool-idle"
    assert row["Recreate pool (rollback)"] == "create ltm pool /Common/pool-idle"
    # vs-idle has no collected profiles in the fixtures: no VS rollback.
    assert row["Recreate virtual servers (rollback)"] == ""


def test_inventory_shows_node_and_pool_verdicts_separately():
    """A dead node in a pool capped by dynamic iRules: node OFFLINE, pool
    MANUAL REVIEW, both visible on the same row."""
    parsed = parse_collection(build_collection())
    parsed.irules["/Common/irule-dyn"] = IRule(
        full_path="/Common/irule-dyn",
        partition="Common",
        name="irule-dyn",
        definition="pool $x",
        has_dynamic_pool_selection=True,
    )
    parsed.virtuals["/Common/vs-web"].irules.append("/Common/irule-dyn")
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    row = next(r for r in tables["inventory"].rows if r[0] == "/Common/node-dead")
    assert row[4] == Verdict.OFFLINE_CANDIDATE
    assert row[17] == Verdict.MANUAL_REVIEW


def test_network_columns_are_appended_to_inventory_and_orphan_nodes():
    _, tables = make_tables()
    inventory = tables["inventory"]
    assert inventory.headers[NETWORK_COLUMNS] == ["ARP MAC", "Network note"]
    assert inventory.verdict_columns == (4, 17)  # unchanged by the appended columns

    member_row = next(
        r for r in inventory.rows if r[0] == "/Common/node-web-1" and r[5] == "/Common/pool-web"
    )
    assert member_row[NETWORK_COLUMNS] == [
        "00:00:5e:00:53:01",
        "in ARP table (MAC 00:00:5e:00:53:01)",
    ]

    orphan_nodes = tables["orphan_nodes"]
    assert orphan_nodes.headers[-2:] == ["ARP MAC", "Network note"]
    assert orphan_nodes.verdict_columns == (4,)
    rows = {row[0]: row for row in orphan_nodes.rows}
    # /26 self-IP: node-dead (10.0.0.50) is local, node-orphan (10.0.0.99) routed.
    assert rows["/Common/node-dead"][-1] == "on local subnet, no ARP entry (idle or down)"
    assert rows["/Common/node-orphan"][-1] == "not directly connected (behind a router)"


def analyze(parsed):
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    return build_tables(parsed, correlation, analysis)


def inventory_row(tables, node_path, pool_path=None):
    return next(
        r
        for r in tables["inventory"].rows
        if r[0] == node_path and (pool_path is None or r[5] == pool_path)
    )


def test_inventory_change_columns_sit_at_x_through_aa():
    _, tables = make_tables()
    headers = tables["inventory"].headers
    assert headers[PING_COLUMNS] == ["Ping (from F5)", "Ping note"]
    assert headers[CHANGE_COLUMNS] == [
        "Remove node from pool",
        "Delete node",
        "Create node (rollback)",
        "Add node back to pool (rollback)",
    ]
    assert len(headers) == 27  # last column is AA
    assert all(len(row) == len(headers) for row in tables["inventory"].rows)


def test_inventory_change_commands_for_offline_node():
    _, tables = make_tables()
    row = inventory_row(tables, "/Common/node-dead", "/Common/pool-dead")
    assert row[4] == Verdict.OFFLINE_CANDIDATE
    assert row[PING_COLUMNS] == ["", ""]
    assert row[CHANGE_COLUMNS] == [
        "modify ltm pool /Common/pool-dead members delete { /Common/node-dead:443 }",
        "delete ltm node /Common/node-dead",
        # Non-default node monitor is restored: it is lost with the node.
        "create ltm node /Common/node-dead address 10.0.0.50 monitor /Common/icmp",
        "modify ltm pool /Common/pool-dead members add { /Common/node-dead:443 }",
    ]


def test_inventory_change_commands_for_orphan_node_skip_pool_columns():
    _, tables = make_tables()
    row = inventory_row(tables, "/Common/node-orphan")
    assert row[4] == Verdict.ORPHAN
    assert row[CHANGE_COLUMNS] == [
        "",
        "delete ltm node /Common/node-orphan",
        "create ltm node /Common/node-orphan address 10.0.0.99 monitor /Common/icmp",
        "",
    ]


def test_inventory_change_columns_blank_for_in_use_node():
    _, tables = make_tables()
    row = inventory_row(tables, "/Common/node-web-1", "/Common/pool-web")
    assert row[4] == Verdict.IN_USE
    assert row[CHANGE_COLUMNS] == ["", "", "", ""]


def test_inventory_change_columns_blank_for_manual_review_node():
    parsed = parse_collection(
        build_collection(denied=[{"partition": "Secret", "endpoint": "/mgmt/tm/ltm/node"}])
    )
    tables = analyze(parsed)
    row = inventory_row(tables, "/Common/node-orphan")
    assert row[4] != Verdict.ORPHAN
    assert row[CHANGE_COLUMNS] == ["", "", "", ""]


def test_inventory_rollback_omits_default_node_monitor_and_keeps_priority_group():
    parsed = parse_collection(build_collection())
    parsed.nodes["/Common/node-dead"].monitor = "default"
    parsed.pools["/Common/pool-dead"].members[0].priority_group = 10
    row = inventory_row(analyze(parsed), "/Common/node-dead", "/Common/pool-dead")
    assert row[CHANGE_COLUMNS][2] == "create ltm node /Common/node-dead address 10.0.0.50"
    assert row[CHANGE_COLUMNS][3] == (
        "modify ltm pool /Common/pool-dead members add "
        "{ /Common/node-dead:443 { priority-group 10 } }"
    )


def test_inventory_change_commands_handle_ipv6_named_and_fqdn_nodes():
    parsed = parse_collection(build_collection())
    parsed.nodes["/Common/2001:db8::10"] = Node(
        full_path="/Common/2001:db8::10",
        partition="Common",
        name="2001:db8::10",
        address="2001:db8::10",
        availability="offline",
    )
    parsed.nodes["/Common/app.example.net"] = Node(
        full_path="/Common/app.example.net",
        partition="Common",
        name="app.example.net",
        address="app.example.net",
    )
    parsed.pools["/Common/pool-dead"].members.append(
        PoolMember(
            node_full_path="/Common/2001:db8::10",
            port="443",
            partition="Common",
            availability="offline",
        )
    )
    tables = analyze(parsed)
    ipv6_row = inventory_row(tables, "/Common/2001:db8::10", "/Common/pool-dead")
    assert ipv6_row[4] == Verdict.OFFLINE_CANDIDATE
    assert ipv6_row[CHANGE_COLUMNS][0] == (
        "modify ltm pool /Common/pool-dead members delete { /Common/2001:db8::10.443 }"
    )
    fqdn_row = inventory_row(tables, "/Common/app.example.net")
    assert fqdn_row[4] == Verdict.ORPHAN
    assert fqdn_row[CHANGE_COLUMNS][2] == (
        "create ltm node /Common/app.example.net fqdn { name app.example.net }"
    )


def test_network_columns_degrade_on_old_raw_cache():
    parsed = parse_collection(build_collection(network=False))
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    row = next(r for r in tables["inventory"].rows if r[0] == "/Common/node-web-1")
    assert row[NETWORK_COLUMNS] == ["", "network data not collected"]
    summary = {str(r[0]): r[1] for r in tables["summary"].rows}
    assert summary["Network data (ARP/self-IP)"] == "not collected (old cache)"


def test_network_note_carries_standby_annotation():
    parsed = parse_collection(build_collection(standby=True))
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    row = next(r for r in tables["inventory"].rows if r[0] == "/Common/node-web-1")
    assert row[NETWORK_COLUMNS][1].endswith("[standby unit: ARP reflects this unit only]")


def test_summary_reports_resumed_collection():
    data = build_collection()
    data.meta["resumed_at"] = ["2026-08-26T10:00:00+00:00"]
    parsed = parse_collection(data)
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    summary = {str(r[0]): r[1] for r in tables["summary"].rows}
    assert summary["Collection resumed (mixed timestamps)"] == "2026-08-26T10:00:00+00:00"


def test_summary_contains_system_info_and_counts():
    _, tables = make_tables()
    rows = {str(row[0]): row[1] for row in tables["summary"].rows}
    assert rows["Hostname"] == "bigip1.example.net"
    assert rows["HA state"] == "active"
    assert Verdict.ORPHAN in rows


def test_write_xlsx(tmp_path):
    _, tables = make_tables()
    out = tmp_path / "report.xlsx"
    write_xlsx(tables, str(out))

    workbook = load_workbook(str(out))
    assert "Summary" in workbook.sheetnames
    assert "Inventory" in workbook.sheetnames

    inventory = workbook["Inventory"]
    assert inventory.freeze_panes == "A2"
    assert inventory.auto_filter.ref is not None
    assert inventory.cell(row=1, column=1).value == "Node"
    assert inventory.max_row > 1

    # Verdict cells carry the conditional fill colors, in both the node
    # verdict and the pool verdict columns.
    node_verdict_column, pool_verdict_column = (c + 1 for c in tables["inventory"].verdict_columns)
    fills = set()
    node_fills = set()
    for row_index in range(2, inventory.max_row + 1):
        cell = inventory.cell(row=row_index, column=pool_verdict_column)
        if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb:
            fills.add(cell.fill.fgColor.rgb)
        node_cell = inventory.cell(row=row_index, column=node_verdict_column)
        if node_cell.fill and node_cell.fill.fgColor and node_cell.fill.fgColor.rgb:
            node_fills.add(node_cell.fill.fgColor.rgb)
    assert "00C6EFCE" in fills or "FFC6EFCE" in fills  # green for IN USE
    assert "00FFC7CE" in node_fills or "FFFFC7CE" in node_fills  # red for ORPHAN node

    dead_chains = workbook["Dead Chains"]
    (verdict_column,) = (c + 1 for c in tables["dead_chains"].verdict_columns)
    cell = dead_chains.cell(row=2, column=verdict_column)
    assert cell.fill.fgColor.rgb in ("00CCC0DA", "FFCCC0DA")  # purple for OFFLINE


def test_write_csv(tmp_path):
    _, tables = make_tables()
    written = write_csv(tables, str(tmp_path))
    assert len(written) == 8

    with open(tmp_path / "inventory.csv", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    assert rows[0][0] == "Node"
    assert len(rows) > 1


def test_default_report_name():
    name = default_report_name("bigip1.example.net")
    assert name.startswith("f5audit_bigip1.example.net_")
    assert name.endswith(".xlsx")
    assert default_report_name("host", "csv").endswith(
        ("0", "1", "2", "3", "4", "5", "6", "7", "8", "9")
    )


def test_iapp_owned_objects_get_no_delete_commands():
    parsed = parse_collection(build_collection())
    for member in parsed.pools["/Common/pool-dead"].members:
        member.app_service = "/Common/adfs.app/adfs"
    correlation = correlate(parsed)
    analysis = Analyzer(parsed, correlation).run()
    tables = build_tables(parsed, correlation, analysis)
    node_row = next(r for r in tables["orphan_nodes"].rows if r[0] == "/Common/node-dead")
    assert node_row[4] == Verdict.MANUAL_REVIEW
    assert node_row[6] == ""
    chain = next(r for r in tables["dead_chains"].rows if r[0] == "/Common/pool-dead")
    assert "delete ltm pool" not in chain[-1]
    assert "delete ltm node" not in chain[-1]
    inventory = inventory_row(tables, "/Common/node-dead", "/Common/pool-dead")
    assert inventory[CHANGE_COLUMNS] == ["", "", "", ""]
